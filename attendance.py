"""
Attendance Detention Tracker
============================

Reads a pasted attendance report, records each tracked absence code in an Excel
workbook, and tells you which students have earned a detention.

Nothing in this program touches the internet. Student data stays on this
computer, in the Excel file you point it at.

HOW IT WORKS, IN ONE PARAGRAPH
------------------------------
The Excel workbook has a sheet called "Log". One row in that sheet = one
recorded absence code for one student for one period on one date. That log is
the single source of truth. Everything else -- the running counts, the totals,
the detention list -- is recalculated from the log every time the program runs.
That means a mistake can always be fixed by deleting a row from the Log sheet
and running the program again.

USAGE
-----
    python attendance.py --workbook attendance.xlsx --input todays_report.txt
    python attendance.py --workbook attendance.xlsx            (then paste, then Ctrl-D)

    --date 2026-09-14   Use a specific date instead of today.
    --dry-run           Show what would change without saving anything.
"""

import argparse
import datetime as dt
import difflib
import re
import sys
from collections import defaultdict
from dataclasses import dataclass

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================================
# SETTINGS -- the things you are most likely to want to change
# ============================================================================

# Which absence codes to count, and how many it takes to earn a detention.
# These are the built-in defaults. If the workbook has a "Settings" sheet,
# that sheet wins and these are ignored.
DEFAULT_TRACKED_CODES = {
    "L": 3,
}

# How an occurrence is counted.
#   "per_period" -- a student absent 4 periods with code L on one day gets 4.
#   "per_day"    -- the same student gets 1, no matter how many periods.
COUNT_MODE = "per_day"

# The pasted report lines up the period columns using a fixed number of
# characters per column. In the sample data each period occupies 5 characters
# ("L-   ", "T-PA "). Change this if your report uses a different width.
PERIOD_FIELD_WIDTH = 5

# The number of the leftmost period column in the report header.
FIRST_PERIOD = 0

# Names in the paste are matched to names in the workbook after being
# simplified this way: trimmed, extra spaces removed, upper-cased.
# Anything closer than this to an existing name, but not an exact match,
# is flagged as a possible typo rather than silently added as a new student.
NEAR_MATCH_SENSITIVITY = 0.90

SHEET_LOG = "Log"
SHEET_STUDENTS = "Students"
SHEET_DETENTIONS = "Detentions"
SHEET_SETTINGS = "Settings"
SHEET_WEEKLY_REPORTS = "Weekly Reports"
SHEET_INTERNAL = "_Internal"   # bookkeeping the program uses; not meant to be edited by hand


# ============================================================================
# PART 1 -- READING THE PASTED REPORT
# ============================================================================

@dataclass
class Occurrence:
    """One absence code, for one student, for one period, on one date."""
    name: str       # as written in the report, e.g. "Anderson, Marcus"
    grade: str      # e.g. "8"
    code: str       # e.g. "L"
    raw_code: str   # e.g. "L-" or "T-PA", kept so nothing is lost
    period: int     # e.g. 0


# A report line looks like:   Anderson, Marcus <TAB> 8 <TAB> X-   X-   X-
# Some exports use runs of spaces instead of tabs, so both are allowed.
LINE_PATTERN = re.compile(
    r"""^\s*
        (?P<name>[A-Za-z][^\t]*?)      # student name, e.g. "Anderson, Marcus"
        (?:\t+|\s{2,})                 # tab(s), or two or more spaces
        (?P<grade>\w{1,3})             # grade, e.g. "8" or "12" or "K"
        (?:\t|\s)                      # one separator before the period columns
        (?P<codes>.*)$                 # everything else is period columns
    """,
    re.VERBOSE,
)

# Lines that are page furniture, not students.
HEADER_HINTS = ("ABS. DATE", "ABS DATE", "GRD", "PAGE ", "TOTAL")


def looks_like_header(line: str) -> bool:
    """True for blank lines, column headings, page numbers and similar."""
    stripped = line.strip()
    if not stripped:
        return True
    upper = stripped.upper()
    return any(hint in upper for hint in HEADER_HINTS)


def split_period_columns(codes_text: str) -> list[tuple[int, str]]:
    """
    Chop the right-hand side of a line into its period columns.

    The columns are fixed width, so position tells us the period number.
    Empty columns are skipped. Returns pairs of (period number, code text).

    >>> split_period_columns("L-")
    [(0, 'L-')]
    >>> split_period_columns("     T-   T-")
    [(1, 'T-'), (2, 'T-')]
    """
    columns = []
    for index in range(0, len(codes_text), PERIOD_FIELD_WIDTH):
        chunk = codes_text[index:index + PERIOD_FIELD_WIDTH].strip()
        if chunk:
            period = FIRST_PERIOD + (index // PERIOD_FIELD_WIDTH)
            columns.append((period, chunk))
    return columns


def base_code(raw_code: str) -> str:
    """
    Reduce a raw entry to the code we track.

    "L-"    -> "L"
    "T-PA"  -> "T"
    "X-"    -> "X"
    """
    return raw_code.split("-", 1)[0].strip().upper()


def parse_pasted_report(text: str) -> tuple[list[Occurrence], list[str]]:
    """
    Turn the pasted text into a list of Occurrences.

    Also returns a list of lines that could not be understood, so nothing is
    ever dropped without the user being told about it.
    """
    occurrences = []
    unparsed = []

    for line in text.splitlines():
        if looks_like_header(line):
            continue

        match = LINE_PATTERN.match(line.rstrip())
        if not match:
            unparsed.append(line.rstrip())
            continue

        name = " ".join(match.group("name").split())
        grade = match.group("grade").strip()

        for period, raw_code in split_period_columns(match.group("codes")):
            occurrences.append(
                Occurrence(
                    name=name,
                    grade=grade,
                    code=base_code(raw_code),
                    raw_code=raw_code,
                    period=period,
                )
            )

    return occurrences, unparsed


# ============================================================================
# PART 2 -- THE COUNTING RULE
# ============================================================================

def keep_only_tracked(occurrences, tracked_codes):
    """Throw away codes we are not counting (X, W, T, M, I and so on)."""
    return [occ for occ in occurrences if occ.code in tracked_codes]


def collapse_to_one_per_day(occurrences):
    """
    Used when COUNT_MODE is "per_day": a student absent five periods with the
    same code counts once, not five times. The earliest period is kept so the
    log still shows where it started.
    """
    best = {}
    for occ in occurrences:
        key = (name_key(occ.name), occ.code)
        if key not in best or occ.period < best[key].period:
            best[key] = occ
    return list(best.values())


def detentions_from_total(total: int, threshold: int) -> tuple[int, int]:
    """
    Work out detentions earned and the count toward the next one.

    The rule is: every time the count reaches the threshold, a detention is
    issued and the count restarts. Over a whole year that is just division:
    7 absences at a threshold of 3 is 2 detentions with 1 left on the clock.

    >>> detentions_from_total(7, 3)
    (2, 1)
    """
    if threshold <= 0:
        return 0, total
    return total // threshold, total % threshold


# ============================================================================
# PART 3 -- NAME MATCHING
# ============================================================================

def name_key(name: str) -> str:
    """
    The simplified form of a name used for matching.

    "  anderson,  marcus " and "Anderson, Marcus" both become "ANDERSON, MARCUS".
    """
    return " ".join(name.split()).upper()


def find_near_matches(new_names, known_names):
    """
    Look for names that are nearly, but not exactly, a name we already have.

    This catches "Halvorsen, Erik" vs "Halvorson, Erik" before it turns into
    two half-counted students.
    """
    warnings = []
    known_keys = [name_key(n) for n in known_names]
    for name in new_names:
        close = difflib.get_close_matches(
            name_key(name), known_keys, n=1, cutoff=NEAR_MATCH_SENSITIVITY
        )
        if close:
            warnings.append((name, close[0]))
    return warnings


# ============================================================================
# PART 4 -- THE WORKBOOK
# ============================================================================

HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
HEADER_FONT = Font(name="Arial", bold=True)
BODY_FONT = Font(name="Arial")
NOTE_FONT = Font(name="Arial", italic=True, color="808080")


def create_template_workbook() -> Workbook:
    """
    Build a brand-new, empty workbook with the four sheets this program
    expects: Settings, Log, Students, Detentions. Students and Detentions
    are rebuilt from the Log every run, so they start out empty here --
    only the Settings sheet needs an example row to be useful right away.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)

    settings = workbook.create_sheet(SHEET_SETTINGS)
    headers = ["Code", "Description", "Track?", "Threshold"]
    write_header_row(settings, headers)
    for code, threshold in DEFAULT_TRACKED_CODES.items():
        settings.cell(2, 1, code).font = BODY_FONT
        settings.cell(2, 3, "YES").font = BODY_FONT
        settings.cell(2, 4, threshold).font = BODY_FONT
    autosize(settings, headers)

    log = workbook.create_sheet(SHEET_LOG)
    headers = ["Date", "Student", "Grade", "Code", "Period", "Raw Code"]
    write_header_row(log, headers)
    autosize(log, headers)

    students = workbook.create_sheet(SHEET_STUDENTS)
    headers = [
        "Student", "Grade", "Code", "Threshold",
        "Total Occurrences", "Detentions Earned", "Count Toward Next",
    ]
    write_header_row(students, headers)
    autosize(students, headers)

    detentions = workbook.create_sheet(SHEET_DETENTIONS)
    headers = ["Date Earned", "Student", "Grade", "Code", "Detention #", "Assigned By", "Served"]
    write_header_row(detentions, headers)
    autosize(detentions, headers)

    return workbook


def read_settings(workbook) -> dict[str, int]:
    """
    Read the tracked codes from the Settings sheet, so that turning a code on
    or changing a threshold never requires editing this program.

    Falls back to DEFAULT_TRACKED_CODES if the sheet is missing.
    """
    if SHEET_SETTINGS not in workbook.sheetnames:
        return dict(DEFAULT_TRACKED_CODES)

    sheet = workbook[SHEET_SETTINGS]
    tracked = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        code = str(row[0]).strip().upper()
        counts = str(row[2]).strip().upper() if len(row) > 2 and row[2] else "NO"
        threshold = row[3] if len(row) > 3 else None
        if counts in ("YES", "Y", "TRUE", "1") and threshold:
            tracked[code] = int(threshold)

    return tracked or dict(DEFAULT_TRACKED_CODES)


def as_date(value):
    """
    Turn whatever is in a date cell into a real date.

    Excel usually hands back a datetime, but a date typed by hand comes back as
    text, and comparing text to a date would crash the program.
    """
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        try:
            return dt.datetime.strptime(value.strip(), "%Y-%m-%d").date()
        except ValueError:
            return None
    return None


def read_log(workbook) -> list[dict]:
    """Read every row already in the Log sheet."""
    sheet = workbook[SHEET_LOG]
    entries = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:          # column B is the name
            continue
        entries.append(
            {
                "date": as_date(row[0]),
                "name": str(row[1]),
                "grade": str(row[2]) if row[2] is not None else "",
                "code": str(row[3]).upper(),
                "period": row[4],
                "raw_code": str(row[5]) if len(row) > 5 and row[5] else "",
            }
        )
    return entries


def already_logged(log_entries, date, code) -> int:
    """How many rows are already recorded for this date and code."""
    return sum(1 for e in log_entries if e["date"] == date and e["code"] == code)


def rebuild_derived_sheets(workbook, log_entries, tracked_codes):
    """
    Recreate the Students and Detentions sheets from the Log.

    The Students sheet uses Excel formulas rather than numbers typed in by this
    program, so that anyone can click a cell and see exactly where the figure
    came from. The program itself never reads those cells back -- it always
    recounts from the Log.
    """
    # ---- roster, in the order names were first seen --------------------
    roster = {}
    for entry in log_entries:
        key = (name_key(entry["name"]), entry["code"])
        if key not in roster:
            roster[key] = {"name": entry["name"], "grade": entry["grade"], "code": entry["code"]}
        elif entry["grade"]:
            roster[key]["grade"] = entry["grade"]

    # ---- Students sheet -------------------------------------------------
    sheet = workbook[SHEET_STUDENTS]
    workbook.remove(sheet)
    sheet = workbook.create_sheet(SHEET_STUDENTS, 0)

    headers = [
        "Student", "Grade", "Code", "Threshold",
        "Total Occurrences", "Detentions Earned", "Count Toward Next",
    ]
    write_header_row(sheet, headers)

    log_rows = max(len(log_entries) + 1, 2)
    for row_number, record in enumerate(sorted(roster.values(), key=lambda r: r["name"]), start=2):
        sheet.cell(row_number, 1, record["name"]).font = BODY_FONT
        sheet.cell(row_number, 2, record["grade"]).font = BODY_FONT
        sheet.cell(row_number, 3, record["code"]).font = BODY_FONT

        # Threshold is looked up from the Settings sheet.
        sheet.cell(row_number, 4,
                   f'=IFERROR(INDEX(Settings!$D$2:$D$50,'
                   f'MATCH(C{row_number},Settings!$A$2:$A$50,0)),"")').font = BODY_FONT

        # Total is counted straight out of the Log sheet.
        sheet.cell(row_number, 5,
                   f'=COUNTIFS(Log!$B$2:$B${log_rows},A{row_number},'
                   f'Log!$D$2:$D${log_rows},C{row_number})').font = BODY_FONT

        # Detentions earned, and what is left on the clock.
        sheet.cell(row_number, 6,
                   f'=IF(D{row_number}="","",ROUNDDOWN(E{row_number}/D{row_number},0))').font = BODY_FONT
        sheet.cell(row_number, 7,
                   f'=IF(D{row_number}="","",E{row_number}-F{row_number}*D{row_number})').font = BODY_FONT

    autosize(sheet, headers)
    sheet.freeze_panes = "A2"

    # ---- Detentions sheet ----------------------------------------------
    # A detention is an event with a date, so these are recorded as plain
    # values: the date is the day the third absence landed.
    sheet = workbook[SHEET_DETENTIONS]
    workbook.remove(sheet)
    sheet = workbook.create_sheet(SHEET_DETENTIONS, 2)

    headers = ["Date Earned", "Student", "Grade", "Code", "Detention #", "Assigned By", "Served"]
    write_header_row(sheet, headers)

    row_number = 2
    for detention in calculate_detentions(log_entries, tracked_codes):
        sheet.cell(row_number, 1, detention["date"]).font = BODY_FONT
        sheet.cell(row_number, 1).number_format = "yyyy-mm-dd"
        sheet.cell(row_number, 2, detention["name"]).font = BODY_FONT
        sheet.cell(row_number, 3, detention["grade"]).font = BODY_FONT
        sheet.cell(row_number, 4, detention["code"]).font = BODY_FONT
        sheet.cell(row_number, 5, detention["number"]).font = BODY_FONT
        row_number += 1

    autosize(sheet, headers)
    sheet.freeze_panes = "A2"


def calculate_detentions(log_entries, tracked_codes) -> list[dict]:
    """
    Walk through the log in date order and note every point at which a student
    crossed the threshold. This is what produces the detention list.
    """
    def sort_key(entry):
        date = as_date(entry["date"]) or dt.date.min
        return (date, entry["name"], entry["period"] or 0)

    running = defaultdict(int)
    issued = defaultdict(int)
    detentions = []

    for entry in sorted(log_entries, key=sort_key):
        threshold = tracked_codes.get(entry["code"])
        if not threshold:
            continue

        key = (name_key(entry["name"]), entry["code"])
        running[key] += 1

        while running[key] >= threshold:
            running[key] -= threshold
            issued[key] += 1
            detentions.append(
                {
                    "date": entry["date"],
                    "name": entry["name"],
                    "grade": entry["grade"],
                    "code": entry["code"],
                    "number": issued[key],
                }
            )

    return detentions


def detentions_since(all_detentions, since_date, through_date) -> list[dict]:
    """
    Filter a detention list down to the ones that are "new" for a weekly
    report: earned after `since_date` (the last time the report was sent,
    or None if it has never been sent) and no later than `through_date`.
    """
    return [
        d for d in all_detentions
        if d["date"] and (since_date is None or d["date"] > since_date)
        and d["date"] <= through_date
    ]


def days_since_last_report(log_entries, last_sent, through_date):
    """
    How many days it has been since the weekly report was last sent, as of
    `through_date`. If no report has ever been sent, this counts from the
    earliest date in the Log instead -- so the warning still fires once
    enough days of un-reported data have piled up. Returns None if there is
    nothing to measure from yet (no log entries and no report history).
    """
    if last_sent is not None:
        reference = last_sent
    else:
        dates = [e["date"] for e in log_entries if e["date"]]
        reference = min(dates) if dates else None

    if reference is None:
        return None
    return (through_date - reference).days


def read_last_report_date(workbook):
    """The date through which the weekly report was last sent, or None."""
    if SHEET_INTERNAL not in workbook.sheetnames:
        return None
    sheet = workbook[SHEET_INTERNAL]
    return as_date(sheet.cell(2, 2).value)


def write_last_report_date(workbook, through_date):
    """Record that the weekly report has now been sent through this date."""
    if SHEET_INTERNAL not in workbook.sheetnames:
        sheet = workbook.create_sheet(SHEET_INTERNAL)
        sheet.sheet_state = "hidden"
        sheet.cell(1, 1, "Key")
        sheet.cell(1, 2, "Value")
        sheet.cell(2, 1, "Weekly report last sent through")
    else:
        sheet = workbook[SHEET_INTERNAL]
    sheet.cell(2, 2, through_date)
    sheet.cell(2, 2).number_format = "yyyy-mm-dd"


WEEKLY_REPORT_HEADERS = [
    "Report Sent On", "Covers After", "Covers Through",
    "Student", "Grade", "Code", "Detention #", "Date Earned",
]


def record_weekly_report(workbook, sent_on, covers_after, covers_through, detentions):
    """
    Permanently save what a weekly report contained, so there's a lasting
    record of exactly who was on each report and when it went out. Rows are
    only ever appended here -- nothing in this sheet is rebuilt or
    recalculated, so a past report can never be changed by anything that
    happens afterward.
    """
    if SHEET_WEEKLY_REPORTS not in workbook.sheetnames:
        sheet = workbook.create_sheet(SHEET_WEEKLY_REPORTS)
        write_header_row(sheet, WEEKLY_REPORT_HEADERS)
    else:
        sheet = workbook[SHEET_WEEKLY_REPORTS]

    row_number = sheet.max_row + 1
    rows = sorted(detentions, key=lambda d: (d["date"], d["name"])) if detentions else [None]

    for detention in rows:
        sheet.cell(row_number, 1, sent_on).number_format = "yyyy-mm-dd"
        if covers_after:
            sheet.cell(row_number, 2, covers_after).number_format = "yyyy-mm-dd"
        sheet.cell(row_number, 3, covers_through).number_format = "yyyy-mm-dd"
        if detention:
            sheet.cell(row_number, 4, detention["name"])
            sheet.cell(row_number, 5, detention["grade"])
            sheet.cell(row_number, 6, detention["code"])
            sheet.cell(row_number, 7, detention["number"])
            sheet.cell(row_number, 8, detention["date"]).number_format = "yyyy-mm-dd"
        else:
            sheet.cell(row_number, 4, "(nobody -- report was empty)")
        for column in range(1, 9):
            sheet.cell(row_number, column).font = BODY_FONT
        row_number += 1

    autosize(sheet, WEEKLY_REPORT_HEADERS)
    sheet.freeze_panes = "A2"


def append_to_log(workbook, occurrences, date):
    """Add today's occurrences to the bottom of the Log sheet."""
    sheet = workbook[SHEET_LOG]
    row_number = sheet.max_row + 1

    # An empty template has a blank row 2 waiting to be used.
    if row_number == 3 and sheet.cell(2, 2).value is None:
        row_number = 2

    for occ in sorted(occurrences, key=lambda o: (o.name, o.period)):
        sheet.cell(row_number, 1, date).number_format = "yyyy-mm-dd"
        sheet.cell(row_number, 2, occ.name)
        sheet.cell(row_number, 3, occ.grade)
        sheet.cell(row_number, 4, occ.code)
        sheet.cell(row_number, 5, occ.period)
        sheet.cell(row_number, 6, occ.raw_code)
        for column in range(1, 7):
            sheet.cell(row_number, column).font = BODY_FONT
        row_number += 1


def write_header_row(sheet, headers):
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(1, column, title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def autosize(sheet, headers):
    widths = [len(str(h)) + 2 for h in headers]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        for index, value in enumerate(row):
            if index < len(widths) and value is not None:
                widths[index] = max(widths[index], min(len(str(value)) + 2, 40))
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


# ============================================================================
# PART 5 -- RUNNING THE PROGRAM
# ============================================================================

def read_input_text(path: str | None) -> str:
    if path:
        with open(path, "r", encoding="utf-8") as handle:
            return handle.read()

    print("Paste the attendance report below.")
    print("When you are finished, press Ctrl-D (Mac/Linux) or Ctrl-Z then Enter (Windows).\n")
    return sys.stdin.read()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", help="the Excel file to update")
    parser.add_argument("--input", help="a text file holding the pasted report")
    parser.add_argument("--date", help="the date of this report, as YYYY-MM-DD (default: today)")
    parser.add_argument("--dry-run", action="store_true", help="show changes without saving")
    parser.add_argument("--force", action="store_true", help="record even if this date is already logged")
    args = parser.parse_args()

    workbook_path = args.workbook
    while not workbook_path:
        workbook_path = input("Path to the Excel workbook: ").strip().strip('"')
    args.workbook = workbook_path

    report_date = (
        dt.datetime.strptime(args.date, "%Y-%m-%d").date() if args.date else dt.date.today()
    )

    workbook = load_workbook(args.workbook)
    tracked_codes = read_settings(workbook)
    existing_log = read_log(workbook)
    last_sent = read_last_report_date(workbook)

    print(f"Report date:   {report_date}")
    print(f"Codes counted: " + ", ".join(f"{c} (every {n})" for c, n in tracked_codes.items()))
    print()

    # ---- read the paste -------------------------------------------------
    occurrences, unparsed = parse_pasted_report(read_input_text(args.input))
    if unparsed:
        print(f"!! {len(unparsed)} line(s) could not be read and were skipped:")
        for line in unparsed[:10]:
            print(f"     {line}")
        print()

    tracked = keep_only_tracked(occurrences, tracked_codes)
    if COUNT_MODE == "per_day":
        tracked = collapse_to_one_per_day(tracked)

    print(f"Lines read:            {len(occurrences)} coded entries")
    print(f"Counting toward rule:  {len(tracked)}")

    if not tracked:
        print("\nNothing to record. No changes made.")
        return

    # ---- guard against pasting the same day twice -----------------------
    for code in tracked_codes:
        duplicates = already_logged(existing_log, report_date, code)
        if duplicates and not args.force:
            print(
                f"\n!! STOPPING: {duplicates} '{code}' entries are already recorded "
                f"for {report_date}.\n"
                f"   Recording again would issue detentions that were not earned.\n"
                f"   Use --force if you are certain you want to add these anyway."
            )
            return

    # ---- warn about names that look like typos --------------------------
    known = {e["name"] for e in existing_log}
    incoming = {o.name for o in tracked}
    brand_new = {n for n in incoming if name_key(n) not in {name_key(k) for k in known}}
    for name, looks_like in find_near_matches(brand_new, known):
        print(f"!! '{name}' is very close to the existing '{looks_like}'. Check for a typo.")

    if brand_new:
        print("\nNew students added to the roster:")
        for name in sorted(brand_new):
            print(f"    {name}")

    # ---- record, recount, report ----------------------------------------
    before = {d["name"] + d["code"] for d in calculate_detentions(existing_log, tracked_codes)}

    append_to_log(workbook, tracked, report_date)
    updated_log = existing_log + [
        {"date": report_date, "name": o.name, "grade": o.grade,
         "code": o.code, "period": o.period, "raw_code": o.raw_code}
        for o in tracked
    ]
    rebuild_derived_sheets(workbook, updated_log, tracked_codes)

    all_detentions = calculate_detentions(updated_log, tracked_codes)
    new_detentions = [d for d in all_detentions if d["date"] == report_date]

    print("\n" + "=" * 58)
    if new_detentions:
        print(f"DETENTIONS EARNED ON {report_date}")
        print("=" * 58)
        for d in new_detentions:
            print(f"  {d['name']:<28} grade {d['grade']:<3} "
                  f"code {d['code']}   detention #{d['number']}")
    else:
        print("No detentions earned today.")
    print("=" * 58)

    since_last_report = detentions_since(all_detentions, last_sent, report_date)
    print()
    if last_sent:
        print(f"Since last weekly report (after {last_sent}, through {report_date}):")
    else:
        print(f"Since last weekly report (none sent yet, through {report_date}):")
    if since_last_report:
        for d in sorted(since_last_report, key=lambda d: (d["date"], d["name"])):
            print(f"  {d['date']}   {d['name']:<28} grade {d['grade']:<3} "
                  f"code {d['code']}   detention #{d['number']}")
    else:
        print("  None yet.")

    days = days_since_last_report(updated_log, last_sent, report_date)
    if days is not None and days > 7:
        if last_sent:
            print(f"\n!! It has been {days} days since the weekly report was last sent "
                  f"(last sent through {last_sent}). Consider sending it soon.")
        else:
            print(f"\n!! It has been {days} days since attendance started being logged, "
                  f"and no weekly report has ever been sent. Consider sending one soon.")

    if args.dry_run:
        print("\nDry run -- the workbook was NOT saved.")
        return

    try:
        workbook.save(args.workbook)
    except PermissionError:
        print(
            f"\n!! COULD NOT SAVE: '{args.workbook}' is open in another program "
            f"(probably Excel).\n"
            f"   Close the file there and run this program again -- nothing has "
            f"been lost."
        )
        sys.exit(1)
    print(f"\nSaved to {args.workbook}")


if __name__ == "__main__":
    main()
